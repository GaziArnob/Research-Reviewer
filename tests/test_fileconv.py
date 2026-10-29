import io
import zipfile

import pymupdf
import pytest
from docx import Document
from PIL import Image

from core.fileconv.convert import CONVERSIONS, ConvertError, run


def _png(color="white", size=(200, 100)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, "PNG")
    return buf.getvalue()


def _docx() -> bytes:
    d = Document()
    d.add_heading("A Title", 0)
    d.add_paragraph("Hello world, this is a paragraph of text for the converter.")
    d.add_paragraph("point one", style="List Bullet")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(1, 1).text = "a", "d"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pdf(text="Plain text on the first page of this document. " * 6) -> bytes:
    doc = pymupdf.open()
    doc.new_page().insert_textbox(pymupdf.Rect(50, 50, 550, 700), text, fontsize=11)
    return doc.tobytes()


def test_images_to_pdf_keeps_order_and_page_count():
    r = run("image-pdf", [("a.png", _png("red")), ("b.png", _png("blue", (100, 300)))])
    doc = pymupdf.open(stream=r.data, filetype="pdf")
    assert len(doc) == 2 and r.filename == "a.pdf" and doc[1].rect.height > doc[1].rect.width  # portrait for the tall image


def test_docx_to_pdf_and_back_keeps_text():
    r = run("docx-pdf", [("doc.docx", _docx())])
    pdf = pymupdf.open(stream=r.data, filetype="pdf")
    text = pdf[0].get_text()
    assert "A Title" in text and "Hello world" in text and "point one" in text
    back = run("pdf-docx", [("doc.pdf", r.data)])
    assert "Hello world" in " ".join(p.text for p in Document(io.BytesIO(back.data)).paragraphs)


def test_pdf_to_text_images_and_docx_to_text():
    pdf = _pdf()
    assert "Plain text" in run("pdf-text", [("x.pdf", pdf)]).data.decode()
    z = zipfile.ZipFile(io.BytesIO(run("pdf-image", [("x.pdf", pdf)]).data))
    assert z.namelist() == ["page_001.png"] and z.read("page_001.png")[:4] == b"\x89PNG"
    txt = run("docx-text", [("d.docx", _docx())]).data.decode()
    assert "Hello world" in txt and "a | " in txt


def test_images_to_docx_embeds_the_picture():
    r = run("image-docx", [("a.png", _png())])
    assert len(Document(io.BytesIO(r.data)).inline_shapes) == 1 and r.filename.endswith(".docx")


def test_wrong_or_damaged_files_give_clear_errors():
    with pytest.raises(ConvertError, match="is not a"):
        run("pdf-docx", [("notes.txt", b"hi")])
    with pytest.raises(ConvertError, match="not a PDF"):
        run("pdf-text", [("x.pdf", b"not really")])
    with pytest.raises(ConvertError, match="valid Word"):
        run("docx-pdf", [("x.docx", b"PK-nope")])
    with pytest.raises(ConvertError, match="could not be read"):
        run("image-pdf", [("x.png", b"garbage")])
    with pytest.raises(ConvertError, match="one file"):
        run("pdf-text", [("a.pdf", _pdf()), ("b.pdf", _pdf())])
    with pytest.raises(ConvertError, match="Choose a conversion"):
        run("nope", [("a.pdf", b"")])
    assert set(CONVERSIONS) >= {"pdf-docx", "docx-pdf", "image-pdf", "image-docx"}
