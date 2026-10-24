import io
import zipfile

import pytest

from core.convert.bundle import ConversionError, convert_source

TEX = r"""
\documentclass{article}
\begin{document}
\title{A Neural Ranking Study}
\author{Ada Lovelace \and Alan Turing}
\maketitle
\begin{abstract}
We study ranking with \emph{neural} models and improve accuracy by 5\% over baselines \cite{smith2020}.
\end{abstract}
\begin{IEEEkeywords}ranking, neural\end{IEEEkeywords}
\section{Introduction}\label{sec:intro}
Ranking matters \cite{smith2020,lee2021}. We describe the method in Section~\ref{sec:method} and show Figure~\ref{fig:a}.
This paragraph is here so that the document has enough words to be converted properly. It talks about retrieval,
ranking, evaluation, datasets, baselines and many other things that papers talk about, repeated a few times so the
converter does not reject the file for being too short. Retrieval ranking evaluation datasets baselines methods results.
More text about the problem setting, the related work and the contributions of this study, written in plain sentences.
\begin{itemize}
\item First contribution with $\alpha$ weights.
\item Second contribution.
\end{itemize}
\section{Method}\label{sec:method}
The score is
\begin{equation}
s(q,d) = \sum_i w_i x_i
\end{equation}
\begin{figure}
\includegraphics[width=\linewidth]{plot}
\caption{Accuracy by model.}\label{fig:a}
\end{figure}
\begin{table}
\caption{Results.}
\begin{tabular}{lc}
\toprule
Model & Acc \\
\bottomrule
\end{tabular}
\end{table}
\bibliographystyle{plain}
\bibliography{refs}
\end{document}
"""
BIB = """@article{smith2020, author={Smith, J.}, title={Ranking models}, journal={J. IR}, year={2020}}
@inproceedings{lee2021, author={Lee, K.}, title={Dense retrieval}, booktitle={SIGIR}, year={2021}}
@article{unused, author={Nobody}, title={Never cited}, year={1999}}
"""
PNG = (b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15\xc4\x89"
       b"\x00\x00\x00\rIDATx\x9cc\xf8\xff\xff?\x00\x05\xfe\x02\xfe\xa7V\x9b\x07\x00\x00\x00\x00IEND\xaeB`\x82")


def _zip(files: dict) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for k, v in files.items():
            z.writestr(k, v)
    return buf.getvalue()


def test_latex_zip_keeps_equations_tables_figures_and_numbers_citations():
    body = TEX.replace("This paragraph is here", "This paragraph is here " * 20)
    c = convert_source(_zip({"main.tex": body, "refs.bib": BIB, "plot.png": PNG}), "paper.zip", "ieee_conference", anonymous=False)
    tex = c.files["main.tex"].decode()
    assert "\\cite{r1,r2}" in tex                                    # numbered in order of first use
    assert "\\begin{equation}" in tex and "\\sum_i w_i x_i" in tex  # equation copied verbatim
    assert "\\toprule" in tex and "\\begin{tabular}" in tex         # table copied, not a placeholder
    assert "images/fig1_plot.png" in tex and "images/fig1_plot.png" in c.files
    assert "Section 2" in tex and "Figure" not in c.warnings        # \ref resolved to the section number
    assert "Ranking models" in tex and "Dense retrieval" in tex and "Never cited" not in tex
    assert "ranking, neural" in tex and "Ada Lovelace" in tex
    assert any("booktabs" in w for w in c.warnings)
    assert "your source file" in c.files["CONVERSION_REPORT.md"].decode()


def test_missing_bib_is_reported_and_bad_input_is_rejected():
    body = TEX.replace("This paragraph is here", "This paragraph is here " * 20)
    c = convert_source(body.encode(), "paper.tex", "neurips2025")
    assert any(".bib" in w for w in c.warnings)
    with pytest.raises(ConversionError):
        convert_source(b"hello", "paper.txt", "neurips2025")
    with pytest.raises(ConversionError):
        convert_source(b"not a zip", "paper.zip", "neurips2025")


def test_zip_paths_cannot_escape_and_docx_headings_are_used():
    from docx import Document
    d = Document()
    d.add_heading("A Word Paper", 0)
    d.add_paragraph("Abstract. We write words " * 5)
    d.add_heading("1 Introduction", 1)
    for _ in range(12):
        d.add_paragraph("A long paragraph of ordinary text about the study and its results. " * 3)
    d.add_paragraph("Item one", style="List Bullet")
    d.add_heading("References", 1)
    d.add_paragraph("[1] Smith. Ranking models. 2020.")
    buf = io.BytesIO()
    d.save(buf)
    c = convert_source(buf.getvalue(), "p.docx", "acl2025")
    tex = c.files["main.tex"].decode()
    assert "\\section{Introduction}" in tex and "\\begin{itemize}" in tex and "Ranking models" in tex
    evil = _zip({"../evil.tex": "x", "main.tex": TEX.replace("This paragraph is here", "This paragraph is here " * 20)})
    assert convert_source(evil, "p.zip", "neurips2025").files["main.tex"]
