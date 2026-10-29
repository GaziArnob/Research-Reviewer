"""General file converter: PDF, Word, images and text. Everything runs on this machine, in memory or in a temp folder
that is deleted at once; nothing is uploaded to another service.

Quality notes (also shown to the user): PDF -> Word keeps layout only approximately; Word -> PDF is a clean re-layout
(LibreOffice is used when installed, otherwise a built-in renderer that handles text, headings, lists, tables and
pictures but not every Word feature); OCR is only as good as the scan.
"""
import io
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

import numpy as np
import pymupdf
from PIL import Image

from core.ingest import ocr

IMAGE_EXT = (".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff", ".webp")
MAX_IMAGES = 30
MAX_PAGES_TO_IMAGES = 100
MAX_PIXELS = 60_000_000
MAX_DOCX_UNZIPPED = 200 * 1024 * 1024
Image.MAX_IMAGE_PIXELS = MAX_PIXELS


class ConvertError(ValueError):
    """A problem the user can understand and fix (wrong file type, damaged file, too large)."""


@dataclass
class Result:
    data: bytes
    filename: str
    media_type: str
    notes: list[str]


# ---------------------------------------------------------------- helpers
def _png(data: bytes) -> bytes:
    """Any supported image -> PNG bytes (also validates it). Transparent images are put on white."""
    try:
        img = Image.open(io.BytesIO(data))
        img.load()
    except Exception as exc:  # noqa: BLE001
        raise ConvertError("One of the images could not be read (damaged or unsupported format).") from exc
    if img.mode in ("RGBA", "LA", "P"):
        img = img.convert("RGBA")
        bg = Image.new("RGB", img.size, "white")
        bg.paste(img, mask=img.split()[-1])
        img = bg
    elif img.mode != "RGB":
        img = img.convert("RGB")
    out = io.BytesIO()
    img.save(out, "PNG")
    return out.getvalue()


def _ocr_image(png: bytes) -> str:
    if not ocr.available():
        raise ConvertError("Text recognition (OCR) is not installed on this server.")
    img = np.asarray(Image.open(io.BytesIO(png)).convert("RGB"))
    result, _ = ocr._get_engine()(img)
    boxes = sorted(result or [], key=lambda r: (round(r[0][0][1] / 12), r[0][0][0]))
    return "\n".join(r[1] for r in boxes)


def _open_pdf(data: bytes) -> pymupdf.Document:
    if not data.startswith(b"%PDF"):
        raise ConvertError("That is not a PDF file.")
    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:  # noqa: BLE001
        raise ConvertError("The PDF could not be opened (damaged file?).") from exc
    if doc.needs_pass:
        raise ConvertError("The PDF is password-protected. Remove the password first.")
    return doc


def _has_text(doc: pymupdf.Document) -> bool:
    return sum(len(p.get_text().strip()) for p in list(doc)[:5]) > 80


def _check_docx(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_DOCX_UNZIPPED or "word/document.xml" not in z.namelist():
                raise ConvertError("That is not a valid Word (.docx) file, or it is too large when unpacked.")
    except zipfile.BadZipFile as exc:
        raise ConvertError("That is not a valid Word (.docx) file.") from exc


def _stem(name: str) -> str:
    return re.sub(r"[^\w\-]", "_", Path(name).stem)[:60] or "file"


# ---------------------------------------------------------------- conversions
def pdf_to_docx(files: list[tuple[str, bytes]], **_) -> Result:
    name, data = files[0]
    doc = _open_pdf(data)
    notes = ["Layout is kept approximately: check tables, columns and equations in Word."]
    from docx import Document
    if not _has_text(doc):
        if not ocr.available():
            raise ConvertError("This PDF has no text layer (a scan) and OCR is not installed.")
        pages = ocr.ocr_pages(data)
        d = Document()
        for i, text in enumerate(pages):
            if i:
                d.add_page_break()
            for line in text.split("\n"):
                d.add_paragraph(line)
        buf = io.BytesIO()
        d.save(buf)
        notes = ["The PDF was a scan, so the text was read with OCR (plain text only, no layout, first 40 pages). Check it for recognition errors."]
        return Result(buf.getvalue(), _stem(name) + ".docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", notes)
    from pdf2docx import Converter
    with tempfile.TemporaryDirectory() as tmp:
        src, dst = Path(tmp) / "in.pdf", Path(tmp) / "out.docx"
        src.write_bytes(data)
        cv = Converter(str(src))
        try:
            cv.convert(str(dst))
        except Exception as exc:  # noqa: BLE001
            raise ConvertError("The PDF could not be converted to Word (unusual structure).") from exc
        finally:
            cv.close()
        return Result(dst.read_bytes(), _stem(name) + ".docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", notes)


_CSS = "body{font-family:sans-serif;font-size:11pt;line-height:1.35} h1{font-size:20pt} h2{font-size:16pt} h3{font-size:13pt} " \
       "table{border-collapse:collapse} td,th{border:1px solid #888;padding:3px} img{max-width:100%}"


def docx_to_pdf(files: list[tuple[str, bytes]], **_) -> Result:
    name, data = files[0]
    _check_docx(data)
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:  # the best result: a real office suite, when one is installed on the server
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "in.docx"
            src.write_bytes(data)
            try:
                subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(src)], check=True, timeout=120,
                               capture_output=True, env={**__import__("os").environ, "HOME": tmp})
                return Result((Path(tmp) / "in.pdf").read_bytes(), _stem(name) + ".pdf", "application/pdf", [])
            except Exception:  # noqa: BLE001 - fall back to the built-in renderer
                pass
    import mammoth
    images: dict[str, bytes] = {}

    def keep_image(image):
        with image.open() as fh:
            blob = fh.read()
        key = f"img{len(images)}.png"
        try:
            images[key] = _png(blob)
        except ConvertError:
            return {"src": ""}
        return {"src": key}

    try:
        html = mammoth.convert_to_html(io.BytesIO(data), convert_image=mammoth.images.img_element(keep_image),
                                       style_map="\n".join(["p[style-name='Title'] => h1:fresh", "p[style-name='Subtitle'] => h2:fresh"])).value
    except Exception as exc:  # noqa: BLE001
        raise ConvertError("The Word file could not be read.") from exc
    arch = pymupdf.Archive()
    for k, v in images.items():
        arch.add(v, k)
    story = pymupdf.Story(html=html or "<p></p>", user_css=_CSS, archive=arch)
    buf = io.BytesIO()
    writer = pymupdf.DocumentWriter(buf)
    mediabox = pymupdf.paper_rect("a4")
    where = mediabox + (56, 56, -56, -56)
    more, pages = 1, 0
    while more and pages < 400:
        dev = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
        pages += 1
    writer.close()
    notes = ["Made with the built-in renderer: text, headings, lists, tables and pictures are kept; headers, footers, text boxes and "
             "exact fonts are not. Install LibreOffice on the server for a closer copy."]
    return Result(buf.getvalue(), _stem(name) + ".pdf", "application/pdf", notes)


def images_to_pdf(files: list[tuple[str, bytes]], **_) -> Result:
    if len(files) > MAX_IMAGES:
        raise ConvertError(f"At most {MAX_IMAGES} images at a time.")
    out = pymupdf.open()
    for _name, data in files:
        png = _png(data)
        w, h = Image.open(io.BytesIO(png)).size
        landscape = w > h
        page = out.new_page(width=842 if landscape else 595, height=595 if landscape else 842)  # A4
        page.insert_image(page.rect + (18, 18, -18, -18), stream=png, keep_proportion=True)
    return Result(out.tobytes(deflate=True), _stem(files[0][0]) + ".pdf", "application/pdf", [f"{len(files)} image(s), one per page, in the order you chose them."])


def images_to_docx(files: list[tuple[str, bytes]], ocr_text: bool = False, **_) -> Result:
    if len(files) > MAX_IMAGES:
        raise ConvertError(f"At most {MAX_IMAGES} images at a time.")
    from docx import Document
    from docx.shared import Cm
    d = Document()
    notes = []
    for i, (_name, data) in enumerate(files):
        png = _png(data)
        if i:
            d.add_page_break()
        if ocr_text:
            for line in _ocr_image(png).split("\n"):
                d.add_paragraph(line)
        else:
            d.add_picture(io.BytesIO(png), width=Cm(16))
    notes.append("Text was read with OCR: check it for recognition errors." if ocr_text else "The pictures are placed as images (not editable text). Tick OCR to get editable text.")
    buf = io.BytesIO()
    d.save(buf)
    return Result(buf.getvalue(), _stem(files[0][0]) + ".docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", notes)


def images_to_text(files: list[tuple[str, bytes]], **_) -> Result:
    if len(files) > MAX_IMAGES:
        raise ConvertError(f"At most {MAX_IMAGES} images at a time.")
    text = "\n\n".join(_ocr_image(_png(d)) for _, d in files)
    return Result(text.encode("utf-8"), _stem(files[0][0]) + ".txt", "text/plain; charset=utf-8", ["Read with OCR: check for recognition errors."])


def pdf_to_images(files: list[tuple[str, bytes]], **_) -> Result:
    name, data = files[0]
    doc = _open_pdf(data)
    notes = []
    if len(doc) > MAX_PAGES_TO_IMAGES:
        notes.append(f"Only the first {MAX_PAGES_TO_IMAGES} of {len(doc)} pages were converted.")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_STORED) as z:
        for i, page in enumerate(list(doc)[:MAX_PAGES_TO_IMAGES], start=1):
            z.writestr(f"page_{i:03d}.png", page.get_pixmap(dpi=150, alpha=False).tobytes("png"))
    return Result(buf.getvalue(), _stem(name) + "_pages.zip", "application/zip", notes)


def pdf_to_text(files: list[tuple[str, bytes]], **_) -> Result:
    name, data = files[0]
    doc = _open_pdf(data)
    notes = []
    if _has_text(doc):
        text = "\n\n".join(p.get_text() for p in doc)
    else:
        if not ocr.available():
            raise ConvertError("This PDF has no text layer (a scan) and OCR is not installed.")
        text = "\n\n".join(ocr.ocr_pages(data))
        notes.append("The PDF was a scan: text read with OCR (first 40 pages). Check for recognition errors.")
    return Result(text.encode("utf-8"), _stem(name) + ".txt", "text/plain; charset=utf-8", notes)


def docx_to_text(files: list[tuple[str, bytes]], **_) -> Result:
    name, data = files[0]
    _check_docx(data)
    from docx import Document
    try:
        d = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ConvertError("The Word file could not be read.") from exc
    lines = [p.text for p in d.paragraphs]
    for t in d.tables:
        lines.append("")
        lines += [" | ".join(c.text.strip() for c in row.cells) for row in t.rows]
    return Result("\n".join(lines).encode("utf-8"), _stem(name) + ".txt", "text/plain; charset=utf-8", [])


# key -> (label, accepted extensions, several files allowed, function, offers OCR option)
CONVERSIONS: dict[str, tuple[str, tuple[str, ...], bool, Callable[..., Result], bool]] = {
    "pdf-docx": ("PDF to Word (.docx)", (".pdf",), False, pdf_to_docx, False),
    "docx-pdf": ("Word (.docx) to PDF", (".docx",), False, docx_to_pdf, False),
    "image-pdf": ("Images to one PDF", IMAGE_EXT, True, images_to_pdf, False),
    "image-docx": ("Images to Word (.docx)", IMAGE_EXT, True, images_to_docx, True),
    "image-text": ("Image to text (OCR)", IMAGE_EXT, True, images_to_text, False),
    "pdf-image": ("PDF to images (PNG, ZIP)", (".pdf",), False, pdf_to_images, False),
    "pdf-text": ("PDF to text (.txt)", (".pdf",), False, pdf_to_text, False),
    "docx-text": ("Word (.docx) to text (.txt)", (".docx",), False, docx_to_text, False),
}


def run(kind: str, files: list[tuple[str, bytes]], ocr_text: bool = False) -> Result:
    if kind not in CONVERSIONS:
        raise ConvertError("Choose a conversion.")
    label, exts, many, fn, _ = CONVERSIONS[kind]
    if not files:
        raise ConvertError("Choose a file.")
    if not many and len(files) > 1:
        raise ConvertError(f"{label} takes one file at a time.")
    for name, _data in files:
        if not name.lower().endswith(exts):
            raise ConvertError(f"“{name}” is not a {' / '.join(e.lstrip('.') for e in exts)} file, which {label} needs.")
    return fn(files, ocr_text=ocr_text)
