import io

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from core.convert.model import Doc
from core.convert.templates import PaperTemplate

PAPER = {"letter": (21.59, 27.94), "a4": (21.0, 29.7)}


def _columns(section, n: int) -> None:
    sect = section._sectPr
    for old in sect.findall(qn("w:cols")):
        sect.remove(old)
    cols = OxmlElement("w:cols")
    cols.set(qn("w:num"), str(n))
    cols.set(qn("w:space"), "425")
    sect.append(cols)


def render_docx(doc: Doc, tpl: PaperTemplate, anonymous: bool) -> bytes:
    """A plain Word version using the venue's page setup where we know it. It is NOT the venue's official Word template."""
    d = Document()
    w, h = PAPER.get(tpl.layout.paper, PAPER["letter"])
    margin = Cm(tpl.layout.margins_cm or 2.0)

    def setup(section):
        section.page_width, section.page_height = Cm(w), Cm(h)
        section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = margin

    setup(d.sections[0])
    normal = d.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(tpl.layout.font_pt)
    normal.paragraph_format.space_after = Pt(4)

    d.add_heading(doc.title or "Title", 0)
    d.add_paragraph("Anonymous Authors" if anonymous else "\n".join(doc.authors_raw[:6]) or "Author names")
    if doc.abstract:
        p = d.add_paragraph()
        p.add_run("Abstract. ").bold = True
        p.add_run(doc.abstract).italic = True
    if doc.keywords:
        p = d.add_paragraph()
        p.add_run("Keywords: ").bold = True
        p.add_run(", ".join(doc.keywords))

    body = d.sections[0]
    if tpl.layout.columns == 2:
        body = d.add_section(WD_SECTION.CONTINUOUS)  # title block stays single column, body is two columns
        setup(body)
        _columns(body, 2)

    for s in doc.sections:
        d.add_heading(s.heading, level=min(s.level, 3))
        for b in s.blocks:
            if b.kind in ("para", "equation"):
                d.add_paragraph(b.text)
            elif b.kind == "bullets":
                for it in b.items:
                    d.add_paragraph(it, style="List Bullet")
            else:
                what = "Figure" if b.kind == "figure" else "Table"
                d.add_paragraph(f"[{what} {b.num}: insert here]").runs[0].italic = True
                d.add_paragraph(f"{what} {b.num}. {b.caption}")
    if doc.references:
        d.add_heading("References", level=1)
        for key, raw in doc.references:
            d.add_paragraph(f"[{key}] {raw}" if key.isdigit() else raw)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
