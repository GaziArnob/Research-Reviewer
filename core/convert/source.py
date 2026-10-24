"""Read an author's own LaTeX or Word source (not a PDF) into the same Doc model the PDF reader produces.

A source file is much more reliable than a PDF: headings, lists, equations, tables and citation keys are explicit, so
nothing has to be guessed from layout. Equations and tables are copied as LaTeX; everything else is flattened to text.
"""
import io
import re
import zipfile
from collections import OrderedDict

from core.convert.model import Block, Doc, Sec

_MAX_FILES = 400
_MAX_UNZIPPED = 80 * 1024 * 1024
_GREEK = {"alpha": "α", "beta": "β", "gamma": "γ", "delta": "δ", "epsilon": "ε", "theta": "θ", "lambda": "λ", "mu": "μ", "pi": "π",
          "sigma": "σ", "tau": "τ", "phi": "φ", "omega": "ω", "times": "×", "leq": "≤", "geq": "≥", "le": "≤", "ge": "≥",
          "cdot": "·", "pm": "±", "approx": "≈", "in": "∈", "rightarrow": "→", "to": "→", "infty": "∞", "sum": "∑", "nabla": "∇",
          "neq": "≠", "ldots": "...", "dots": "..."}
_SECTION = {"section": 1, "subsection": 2, "subsubsection": 3}
_STRUCT = re.compile(
    r"\\(?P<sec>section|subsection|subsubsection)\*?\s*(?:\[[^\]]*\])?\s*\{"
    r"|\\begin\{(?P<env>itemize|enumerate|figure\*?|table\*?|equation\*?|align\*?|eqnarray\*?|gather\*?|multline\*?)\}")


def _group(text: str, i: int) -> tuple[str, int]:
    """text[i] is '{'; returns (content, index after the matching '}')."""
    depth, j = 0, i
    while j < len(text):
        c = text[j]
        if c == "\\":
            j += 2
            continue
        depth += c == "{"
        depth -= c == "}"
        j += 1
        if depth == 0:
            return text[i + 1:j - 1], j
    return text[i + 1:], len(text)


def _arg(text: str, command: str) -> str:
    m = re.search(r"\\" + command + r"\*?\s*(?:\[[^\]]*\])?\s*\{", text)
    return _group(text, m.end() - 1)[0] if m else ""


def _drop_commands(text: str, names: tuple) -> str:
    """Remove \\name{...} (with its balanced braces) so title-block commands do not end up as body text."""
    pat = re.compile(r"\\(?:" + "|".join(names) + r")\*?\s*(?:\[[^\]]*\])?\s*\{")
    while (m := pat.search(text)):
        text = text[:m.start()] + text[_group(text, m.end() - 1)[1]:]
    return text


def _strip_comments(text: str) -> str:
    return re.sub(r"(?<!\\)%.*", "", text)


def _math(s: str) -> str:
    s = re.sub(r"\\([A-Za-z]+)", lambda m: _GREEK.get(m.group(1), ""), s)
    return re.sub(r"[{}]", "", s)


class _Cites:
    """Numbers citation keys in order of first use, as BibTeX would for a numeric style."""

    def __init__(self):
        self.order: "OrderedDict[str, int]" = OrderedDict()
        self.labels: dict[str, str] = {}
        self.unresolved = 0

    def cite(self, m: re.Match) -> str:
        nums = []
        for key in (k.strip() for k in m.group(1).split(",") if k.strip()):
            nums.append(self.order.setdefault(key, len(self.order) + 1))
        return "[" + ", ".join(str(n) for n in nums) + "]"

    def ref(self, m: re.Match) -> str:
        if m.group(1) in self.labels:
            return self.labels[m.group(1)]
        self.unresolved += 1
        return "??"


def clean(s: str, c: _Cites, stats: dict) -> str:
    """Flatten inline LaTeX to plain text. Citations become [n]; \\ref becomes the number of its label."""
    s = re.sub(r"\\(?:cite[a-z]*|parencite|textcite|autocite)\*?\s*(?:\[[^\]]*\])*\s*\{([^{}]*)\}", c.cite, s)
    s = re.sub(r"\\(?:ref|autoref|cref|Cref|eqref|pageref)\*?\s*\{([^{}]*)\}", c.ref, s)
    s = re.sub(r"\\footnote\s*\{([^{}]*)\}", r" (\1)", s)
    s = re.sub(r"\\href\s*\{[^{}]*\}\s*\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\url\s*\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\$\$?([^$]*)\$\$?", lambda m: (stats.update(inline_math=stats.get("inline_math", 0) + 1) or _math(m.group(1))), s)
    s = re.sub(r"\\\(([^)]*?)\\\)", lambda m: _math(m.group(1)), s)
    for _ in range(4):
        s = re.sub(r"\\(?:textbf|textit|emph|texttt|textsc|textrm|underline|mbox|text|mathrm|mathbf|mathit|mathcal|bm)\s*\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\(?:label|vspace|hspace|noindent|centering|newpage|clearpage|small|footnotesize|large|Large|bibliographystyle|maketitle)\*?(?:\{[^{}]*\})?", "", s)
    s = s.replace("\\\\", " ").replace("\\,", " ").replace("\\ ", " ").replace("~", " ")
    s = re.sub(r"\\([&%$#_{}])", r"\1", s)
    s = re.sub(r"\\[A-Za-z]+\*?", "", s)
    s = re.sub(r"[{}]", "", s)
    return re.sub(r"[ \t]+", " ", s).strip()


def _paragraphs(text: str, c: _Cites, stats: dict) -> list[Block]:
    out = []
    for chunk in re.split(r"\n\s*\n", text):
        t = clean(" ".join(chunk.split()), c, stats)
        if len(t) > 1:
            out.append(Block("para", text=t))
    return out


def _bib_from_bibtex(bib: str) -> dict[str, str]:
    refs = {}
    for m in re.finditer(r"@(\w+)\s*\{\s*([^,\s]+)\s*,", bib):
        if m.group(1).lower() in ("string", "comment", "preamble"):
            continue
        body = bib[m.end():]
        end = re.search(r"\n\s*@\w+\s*\{", body)
        body = body[:end.start()] if end else body
        fields = {}
        for f in re.finditer(r"(\w+)\s*=\s*(\{|\")", body):
            if f.group(2) == "{":
                val, _ = _group(body, f.end() - 1)
            else:
                val = body[f.end():body.find('"', f.end())]
            fields[f.group(1).lower()] = " ".join(val.split())
        venue = fields.get("journal") or fields.get("booktitle") or fields.get("publisher") or ""
        parts = [fields.get("author", "").replace(" and ", ", "), fields.get("title", ""), venue, fields.get("year", "")]
        refs[m.group(2)] = re.sub(r"[{}]", "", ". ".join(p for p in parts if p)).strip()
    return refs


def _read_files(data: bytes, filename: str) -> dict[str, bytes]:
    if not filename.lower().endswith(".zip"):
        return {filename: data}
    files = {}
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        infos = z.infolist()
        if len(infos) > _MAX_FILES or sum(i.file_size for i in infos) > _MAX_UNZIPPED:
            raise ValueError("The ZIP is too large or has too many files.")
        for i in infos:
            name = i.filename.replace("\\", "/")
            if i.is_dir() or name.startswith("/") or ".." in name.split("/"):
                continue  # never trust paths from an upload
            files[name] = z.read(i)
    return files


def _expand_inputs(text: str, files: dict[str, bytes], depth: int = 0) -> str:
    def repl(m: re.Match) -> str:
        want = m.group(1).strip().replace("\\", "/")
        for cand in (want, want + ".tex"):
            for name, data in files.items():
                if name == cand or name.endswith("/" + cand):
                    body = _strip_comments(data.decode("utf-8", "replace"))
                    return _expand_inputs(body, files, depth + 1) if depth < 3 else body
        return ""
    return re.sub(r"\\(?:input|include)\s*\{([^{}]+)\}", repl, text)


def doc_from_latex(data: bytes, filename: str) -> Doc:
    files = _read_files(data, filename)
    mains = [n for n, d in files.items() if n.lower().endswith(".tex") and b"\\begin{document}" in d]
    if not mains:
        raise ValueError("No .tex file with \\begin{document} was found.")
    main = min(mains, key=lambda n: (n.count("/"), len(n)))
    text = _expand_inputs(_strip_comments(files[main].decode("utf-8", "replace")), files)
    doc, stats, c = Doc(), {}, _Cites()
    doc.title = clean(_arg(text, "title"), c, stats)
    author = _arg(text, "author")
    doc.authors_raw = [a for a in (clean(x, c, stats) for x in re.split(r"\\and|\\\\|\\IEEEauthorblock[NA]", author)) if a][:12]
    abs_m = re.search(r"\\begin\{abstract\}(.*?)\\end\{abstract\}", text, re.S)
    doc.abstract = clean(" ".join(abs_m.group(1).split()), c, stats) if abs_m else ""
    kw = re.search(r"\\begin\{(?:IEEEkeywords|keywords)\}(.*?)\\end\{(?:IEEEkeywords|keywords)\}", text, re.S) or \
        re.search(r"\\keywords\s*\{([^{}]*)\}", text)
    if kw:
        doc.keywords = [k.strip() for k in re.split(r"[,;]", clean(kw.group(1), c, stats)) if k.strip()]

    body = text.split("\\begin{document}", 1)[1]
    bib_m = re.search(r"\\begin\{thebibliography\}.*?\\end\{thebibliography\}", body, re.S)
    bib_cmd = re.search(r"\\(?:bibliography|addbibresource)\s*\{([^{}]*)\}", body)
    cut = min([m.start() for m in (bib_m, re.search(r"\\bibliography\s*\{|\\printbibliography|\\end\{document\}", body)) if m] or [len(body)])
    main_text = body[:cut]
    if abs_m:
        main_text = main_text.replace(abs_m.group(0), "")
    main_text = re.sub(r"\\begin\{(?:IEEEkeywords|keywords)\}.*?\\end\{(?:IEEEkeywords|keywords)\}", "", main_text, flags=re.S)
    main_text = _drop_commands(main_text, ("title", "author", "keywords", "affiliation", "date", "thanks", "address", "email"))

    # Pass 1: label -> number for sections, figures and tables, so \ref can be resolved.
    secs = [0, 0, 0]
    counters = {"figure": 0, "table": 0}
    for m in _STRUCT.finditer(main_text):
        if m.group("sec"):
            lvl = _SECTION[m.group("sec")]
            if not main_text[m.start():m.end()].count("*"):
                secs[lvl - 1] += 1
                for k in range(lvl, 3):
                    secs[k] = 0
            end = _group(main_text, m.end() - 1)[1]
            lab = re.match(r"\s*\\label\{([^{}]*)\}", main_text[end:])
            if lab:
                c.labels[lab.group(1)] = ".".join(str(x) for x in secs[:lvl])
        elif m.group("env").rstrip("*") in counters:
            kind = m.group("env").rstrip("*")
            counters[kind] += 1
            env_end = main_text.find("\\end{" + m.group("env") + "}", m.end())
            for lab in re.finditer(r"\\label\{([^{}]*)\}", main_text[m.end():env_end]):
                c.labels[lab.group(1)] = str(counters[kind])

    # Pass 2: build blocks.
    current = Sec(heading="Introduction", level=1)
    doc.sections.append(current)
    pos, fig_n, tab_n, flattened, unrecovered = 0, 0, 0, 0, 0
    first_heading = True

    def add(blocks):
        current.blocks.extend(blocks)

    for m in _STRUCT.finditer(main_text):
        if m.start() < pos:
            continue
        add(_paragraphs(main_text[pos:m.start()], c, stats))
        if m.group("sec"):
            title, end = _group(main_text, m.end() - 1)
            if first_heading and not current.blocks:
                doc.sections.pop()
            first_heading = False
            current = Sec(heading=clean(title, c, stats), level=_SECTION[m.group("sec")])
            doc.sections.append(current)
            pos = end
            continue
        env = m.group("env")
        end_tag = "\\end{" + env + "}"
        end = main_text.find(end_tag, m.end())
        inner = main_text[m.end():end if end >= 0 else len(main_text)]
        pos = (end + len(end_tag)) if end >= 0 else len(main_text)
        kind = env.rstrip("*")
        if kind in ("itemize", "enumerate"):
            items = [clean(" ".join(i.split()), c, stats) for i in re.split(r"\\item\b", inner)[1:]]
            add([Block("bullets", items=[i for i in items if i])])
        elif kind == "figure":
            fig_n += 1
            cap = clean(" ".join(_arg(inner, "caption").split()), c, stats)
            blk = Block("figure", num=fig_n, caption=cap)
            g = re.search(r"\\includegraphics\s*(?:\[[^\]]*\])?\s*\{([^{}]+)\}", inner)
            if g:
                base = g.group(1).strip().rsplit("/", 1)[-1]
                for name, img in files.items():
                    stem = name.rsplit("/", 1)[-1]
                    if stem == base or stem.rsplit(".", 1)[0] == base.rsplit(".", 1)[0]:
                        ext = stem.rsplit(".", 1)[-1].lower()
                        if ext in ("png", "jpg", "jpeg"):
                            blk.image = f"fig{fig_n}_{stem}"
                            doc.images[blk.image] = img
                        break
            unrecovered += blk.image is None
            add([blk])
        elif kind == "table":
            tab_n += 1
            tab = re.search(r"\\begin\{tabular\*?\}.*?\\end\{tabular\*?\}", inner, re.S)
            add([Block("table", num=tab_n, caption=clean(" ".join(_arg(inner, "caption").split()), c, stats),
                       raw=tab.group(0) if tab else "")])
        else:  # display equations: keep the LaTeX verbatim
            flattened += 1
            add([Block("equation", text=clean(_math(inner), c, stats), raw=f"\\begin{{{env}}}{inner}\\end{{{env}}}")])
    add(_paragraphs(main_text[pos:], c, stats))
    doc.sections = [s for s in doc.sections if s.blocks or s.level > 1 or s.heading != "Introduction"]

    # References: the .bib entries or \bibitem list, numbered by first citation.
    entries: dict[str, str] = {}
    if bib_m:
        for b in re.finditer(r"\\bibitem\s*(?:\[[^\]]*\])?\s*\{([^{}]*)\}(.*?)(?=\\bibitem|\Z)", bib_m.group(0).split("\\begin{thebibliography}", 1)[1], re.S):
            entries[b.group(1)] = clean(" ".join(b.group(2).replace("\\end{thebibliography}", "").split()), _Cites(), {})
    for name, d in files.items():
        if name.lower().endswith(".bib"):
            entries.update(_bib_from_bibtex(d.decode("utf-8", "replace")))
    missing = [k for k in c.order if k not in entries]
    doc.references = [(str(n), entries.get(k, f"[missing entry: {k}]")) for k, n in c.order.items()]
    if bib_cmd and not any(n.lower().endswith(".bib") for n in files):
        doc.warnings.append("The bibliography file (.bib) was not in the upload: references are placeholders. Upload a ZIP that includes it.")
    if missing and not (bib_cmd and not any(n.lower().endswith(".bib") for n in files)):
        doc.warnings.append(f"{len(missing)} cited key(s) have no bibliography entry: {', '.join(missing[:5])}.")
    if c.unresolved:
        doc.warnings.append(f"{c.unresolved} cross-reference(s) (equations, appendices, other labels) could not be resolved and show as ??.")
    if flattened:
        doc.warnings.append(f"{flattened} display equation(s) were copied verbatim. Packages they need (amsmath etc.) must be in the venue preamble.")
    if stats.get("inline_math"):
        doc.warnings.append(f"{stats['inline_math']} inline formula(s) were flattened to plain text: check subscripts and symbols.")
    if unrecovered:
        doc.warnings.append(f"{unrecovered} figure(s) have no PNG/JPEG image in the upload (PDF/EPS figures are not copied) and are placeholders.")
    if any("\\toprule" in b.raw for s in doc.sections for b in s.blocks if b.kind == "table"):
        doc.warnings.append("Tables were copied as LaTeX and use booktabs: add \\usepackage{booktabs} if the venue style lacks it.")
    doc.stats = {"source": "latex", "main_file": main}
    return doc


def doc_from_docx(data: bytes) -> Doc:
    from docx import Document
    try:
        d = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ValueError("That is not a readable .docx file.") from exc
    doc, c, stats = Doc(), _Cites(), {}
    current, mode = None, "front"
    refs: list[str] = []
    fig = tab = 0
    for p in d.paragraphs:
        text = " ".join(p.text.split())
        if not text:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style == "title" and not doc.title:
            doc.title = text
        elif style.startswith("heading"):
            lvl = int(re.sub(r"\D", "", style) or 1)
            name = re.sub(r"^\d+(?:\.\d+)*\.?\s+", "", text)
            if name.lower() == "abstract":
                mode = "abstract"
            elif name.lower() in ("references", "bibliography"):
                mode = "refs"
            else:
                mode, current = "body", Sec(heading=name, level=min(3, lvl))
                doc.sections.append(current)
        elif mode == "front":
            if re.match(r"(?i)^(abstract)[\s:.\u2014-]", text):
                doc.abstract = re.sub(r"(?i)^abstract[\s:.\u2014-]*", "", text)
            elif re.match(r"(?i)^(keywords?|index terms)\s*[:\u2014-]", text):
                doc.keywords = [k.strip() for k in re.split(r"[,;]", re.sub(r"(?i)^[^:\u2014-]*[:\u2014-]", "", text, count=1)) if k.strip()]
            elif not doc.title:
                doc.title = text
            elif len(text) < 120:
                doc.authors_raw.append(text)
        elif mode == "abstract":
            doc.abstract = (doc.abstract + " " + text).strip()
        elif mode == "refs":
            refs.append(re.sub(r"^\[?\d+\]?[.)]?\s*", "", text))
        elif current is not None:
            if "list" in style:
                if current.blocks and current.blocks[-1].kind == "bullets":
                    current.blocks[-1].items.append(text)
                else:
                    current.blocks.append(Block("bullets", items=[text]))
            elif (m := re.match(r"^(Figure|Fig\.|Table)\s*(\d+)\s*[:.]\s*(.*)$", text, re.I)):
                kind = "table" if m.group(1).lower() == "table" else "figure"
                current.blocks.append(Block(kind, num=int(m.group(2)), caption=m.group(3)))
                fig, tab = fig + (kind == "figure"), tab + (kind == "table")
            else:
                current.blocks.append(Block("para", text=text))
    if not doc.sections:
        raise ValueError("No headings were found. Use Word's Heading 1 / Heading 2 styles for section titles.")
    doc.references = [(str(i), r) for i, r in enumerate(refs, start=1)]
    if not doc.abstract:
        doc.warnings.append("No abstract was found: add one by hand (a heading called Abstract, or a paragraph starting with 'Abstract').")
    if not refs:
        doc.warnings.append("No reference list was found (a heading called References).")
    if d.tables:
        doc.warnings.append(f"{len(d.tables)} Word table(s) are not copied; only their captions are placeholders.")
    if d.inline_shapes:
        doc.warnings.append(f"{len(d.inline_shapes)} picture(s) in the Word file are not copied; figures are placeholders.")
    doc.stats = {"source": "docx"}
    return doc
