import io
from typing import Optional

from pydantic import BaseModel, Field

from core.models import Finding, ParsedPaper, PipelineResult
from core.templates.autochecks import run_check
from core.templates.loader import Template

DISCLAIMER = ("Prepared by Research Reviewer as evidence for a human reviewer. Scores, ratings, confidence and the "
              "recommendation are deliberately left blank: on a benchmark the tool's own scores did not track human "
              "reviewers. Quotes were checked against the paper text; keyword 'hints' were not. Review the evidence "
              "yourself and write your review in your own words.")


class FormField(BaseModel):
    key: str
    label: str
    type: str
    guidance: str = ""
    source: str = "official"
    scale: Optional[dict] = None
    choices: list[str] = Field(default_factory=list)
    prepared: list[str] = Field(default_factory=list)  # evidence lines prepared by the tool
    human_fill: bool = False


class ChecklistRow(BaseModel):
    id: str
    group: str
    text: str
    options: list[str]
    answer: str  # yes | partial | no | NA | unclear | (empty for manual items)
    note: str = ""
    quote: str = ""
    auto: bool = False


class Form(BaseModel):
    template_id: str
    name: str
    paper_title: str
    provenance_level: str
    provenance_notes: str
    official_urls: list[str]
    criteria: list[str] = Field(default_factory=list)
    fields: list[FormField]
    checklist: list[ChecklistRow] = Field(default_factory=list)
    llm_hidden: bool = False
    llm_policy: str = "not_stated"
    third_party: str = "not_stated"
    policy_notes: list[str] = Field(default_factory=list)


def _line(f: Finding) -> list[str]:
    tag = {"verified": "quote verified", "needs_verification": "verify", "uncertain": "uncertain"}[f.status]
    out = [f"- {f.text} [{tag}]"]
    for e in f.evidence[:1]:
        where = f"p.{e.page}" if e.page else (e.section or "")
        out.append(f'    > "{e.quote[:220]}"{f" ({where})" if where else ""}')
    return out


def _lines(findings: list[Finding], empty: str) -> list[str]:
    return [l for f in findings for l in _line(f)] or [empty]


def _pick(result: PipelineResult, kinds=None, modules=None) -> list[Finding]:
    return [f for f in result.findings if (kinds is None or f.kind in kinds) and (modules is None or f.module in modules)]


NEED_LLM = "(Run the LLM review to prepare this part.)"


def _prepare(fill: str, result: PipelineResult, paper: Optional[ParsedPaper]) -> list[str]:
    weak_modules = ("review", "critic", "novelty", "literature")
    if fill == "summary":
        u = result.understanding
        if not u:
            return [NEED_LLM]
        return [u["summary"]] + [f"- Claimed contribution: {c}" for c in u.get("contributions", [])]
    if fill == "strengths":
        return _lines(_pick(result, {"strength"}), NEED_LLM)
    if fill == "weaknesses":
        return _lines(_pick(result, {"major", "minor"}, weak_modules), "(No evidence-linked weakness was found by the tool.)")
    if fill == "strengths_weaknesses":
        return (["Strengths:"] + _prepare("strengths", result, paper) + ["Weaknesses:"] + _prepare("weaknesses", result, paper))
    if fill == "questions":
        return _lines(_pick(result, {"question"}), NEED_LLM)
    if fill == "claims_evidence":
        claims = [f'- Claim: {c.text}\n    > "{c.quote[:220]}" (p.{c.page})' for c in result.claims]
        return (claims or [NEED_LLM]) + ["Checks on the evidence:"] + _lines(
            _pick(result, None, ("statistics", "critic", "reproducibility")), "(No evidence issue found by keyword checks.)")
    if fill == "prior_work":
        rel = [f"- {p.title} ({p.year or 'n/a'}){f' {p.url}' if p.url else ''}" for p in result.related[:8]]
        return (_lines(_pick(result, None, ("novelty", "literature")), "(No overlap or missing-citation issue found.)")
                + (["Related papers found by search:"] + rel if rel else ["(No related-paper search results: run with online checks.)"]))
    if fill == "other_aspects":
        return _lines(_pick(result, None, ("writing", "preflight")), "(No writing or format issue found.)")
    if fill == "quality":
        return _lines(_pick(result, None, ("statistics", "reproducibility", "references", "citations")), "(No issue found by checks.)")
    if fill == "clarity":
        return _lines(_pick(result, None, ("writing", "preflight")), "(No writing or format issue found.)")
    if fill == "reproducibility":
        return _lines(_pick(result, None, ("reproducibility",)), "(No reproducibility gap found by keyword checks.)")
    if fill == "limitations" and paper is not None:
        ans, ev, note = run_check("limitations", paper, result)
        return [f"Keyword check: {note}" + (f' "{ev.quote[:200]}"' if ev else "")]
    return []


def _checklist(template: Template, paper: Optional[ParsedPaper], result: PipelineResult) -> list[ChecklistRow]:
    rows: list[ChecklistRow] = []
    gates: dict[str, str] = {}
    for item in template.checklist:
        answer, note, quote = "", "", ""
        if item.requires and gates.get(item.requires) == "no":
            answer, note = "NA", "Not applicable: the gate question was answered 'no'."
        elif item.auto and paper is not None:
            answer, ev, note = run_check(item.auto, paper, result)
            quote = ev.quote[:240] if ev else ""
            if item.options and answer not in item.options and answer != "unclear":
                answer = "partial" if "partial" in item.options and answer != "no" else ("no" if "no" in item.options else answer)
        if item.gate:
            gates[item.gate] = answer
        rows.append(ChecklistRow(id=item.id, group=item.group, text=item.text, options=item.options, answer=answer,
                                 note=note, quote=quote, auto=bool(item.auto and paper is not None)))
    return rows


LLM_HIDDEN = "(Not prepared: this venue forbids LLM-written review text. Only rule-based checks are shown.)"


def without_llm(result: PipelineResult) -> PipelineResult:
    """A copy of the result with everything an LLM wrote removed (summary, claims, comments, scores)."""
    return result.model_copy(update={
        "understanding": None, "claims": [], "scores": None, "llm_model": None,
        "findings": [f for f in result.findings if not str(f.source).startswith("llm")]})


def build_form(template: Template, result: PipelineResult, paper: Optional[ParsedPaper]) -> Form:
    hidden = template.policy.llm_review_text == "forbidden"
    if hidden:
        result = without_llm(result)
    fields = []
    for s in template.sections:
        prepared = _prepare(s.fill, result, paper) if s.fill else []
        if hidden:
            prepared = [LLM_HIDDEN if line == NEED_LLM else line for line in prepared]
        fields.append(FormField(key=s.key, label=s.label, type=s.type, guidance=s.guidance, source=s.source,
                                scale=s.scale, choices=s.choices, prepared=prepared,
                                human_fill=s.fill is None or s.type in ("scale", "choice", "ack")))
    return Form(template_id=template.id, name=template.name, paper_title=result.title or result.filename,
                provenance_level=template.provenance.level, provenance_notes=template.provenance.notes,
                official_urls=template.provenance.official_urls, criteria=template.criteria, fields=fields,
                checklist=_checklist(template, paper, result), llm_hidden=hidden,
                llm_policy=template.policy.llm_review_text, third_party=template.policy.third_party_services,
                policy_notes=template.policy.notes)


def _scale_text(scale: Optional[dict]) -> str:
    if not scale:
        return ""
    labels = scale.get("labels") or {}
    parts = [f"{k} = {v}" for k, v in sorted(labels.items(), key=lambda kv: -float(kv[0]))]
    base = f"{scale['min']} to {scale['max']}" + (f" in steps of {scale['step']:g}" if scale.get("step", 1) != 1 else "")
    return base + (": " + "; ".join(parts) if parts else "")


def to_markdown(form: Form) -> str:
    L = [f"# {form.name}", f"Paper: {form.paper_title}", "", f"> {DISCLAIMER}", ""]
    L.append(f"Template check: **{form.provenance_level.upper()}**. {form.provenance_notes}")
    L += [f"Source: {u}" for u in form.official_urls] + [""]
    if form.policy_notes:
        L += [f"## Venue policy on AI tools ({form.llm_policy.replace('_', ' ')})"] + [f"- {n}" for n in form.policy_notes] + [""]
    if form.criteria:
        L += ["## Official evaluation criteria"] + [f"- {c}" for c in form.criteria] + [""]
    for f in form.fields:
        flag = "" if f.source == "official" else f" *(field/scale {f.source}: check the live form)*"
        L += [f"## {f.label}{flag}"]
        if f.guidance:
            L.append(f"*{f.guidance}*")
        if f.type == "scale":
            L.append(f"Scale: {_scale_text(f.scale) or 'see live form'}")
        if f.type == "choice":
            L.append("Options: " + " / ".join(f.choices))
        L += f.prepared
        if f.human_fill:
            L.append("**[ ] To be filled in by the human reviewer**")
        L.append("")
    if form.checklist:
        L += ["## Checklist (keyword hints marked 'auto'; verify each)", "| Item | Answer | Note |", "|---|---|---|"]
        for r in form.checklist:
            ans = (r.answer or "manual") + (" (auto)" if r.auto else "")
            item = (r.group + ": " if r.group else "") + r.text
            note = (r.note + (f' "{r.quote}"' if r.quote else "")).replace("|", "/").replace("\n", " ")
            L.append(f"| {item} | {ans} | {note} |")
    return "\n".join(L)


def to_docx(form: Form) -> bytes:
    from docx import Document
    doc = Document()
    doc.add_heading(form.name, 0)
    doc.add_paragraph(f"Paper: {form.paper_title}")
    doc.add_paragraph(DISCLAIMER)
    doc.add_paragraph(f"Template check: {form.provenance_level.upper()}. {form.provenance_notes}")
    if form.policy_notes:
        doc.add_heading(f"Venue policy on AI tools ({form.llm_policy.replace('_', ' ')})", 1)
        for n in form.policy_notes:
            doc.add_paragraph(n, style="List Bullet")
    if form.criteria:
        doc.add_heading("Official evaluation criteria", 1)
        for c in form.criteria:
            doc.add_paragraph(c, style="List Bullet")
    for f in form.fields:
        doc.add_heading(f.label + ("" if f.source == "official" else f" ({f.source}: check the live form)"), 1)
        if f.guidance:
            doc.add_paragraph().add_run(f.guidance).italic = True
        if f.type == "scale":
            doc.add_paragraph(f"Scale: {_scale_text(f.scale) or 'see live form'}")
        if f.type == "choice":
            doc.add_paragraph("Options: " + " / ".join(f.choices))
        for line in f.prepared:
            doc.add_paragraph(line.strip(), style="List Bullet" if line.startswith("- ") else None)
        if f.human_fill:
            doc.add_paragraph("[ ] To be filled in by the human reviewer")
    if form.checklist:
        doc.add_heading("Checklist", 1)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, ("Item", "Answer", "Note")):
            cell.text = text
        for r in form.checklist:
            cells = table.add_row().cells
            cells[0].text = (r.group + ": " if r.group else "") + r.text
            cells[1].text = (r.answer or "manual") + (" (auto)" if r.auto else "")
            cells[2].text = (r.note + (f' "{r.quote}"' if r.quote else "")).strip()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
